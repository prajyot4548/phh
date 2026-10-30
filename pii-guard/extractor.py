"""
Read plain text out of the file types PII Guard understands.
Optional formats (xlsx, pptx, rtf) degrade gracefully if their library isn't installed.
"""
import os
import config
from logger_setup import get_logger

log = get_logger("extractor")

TEXT_EXT = {".txt", ".csv", ".md", ".json", ".log", ".xml", ".html", ".htm"}
IMG_EXT = {".jpg", ".jpeg", ".png", ".bmp", ".tiff", ".tif", ".webp"}
OFFICE_EXT = {".pdf", ".docx", ".xlsx", ".pptx", ".rtf"}
SUPPORTED = TEXT_EXT | IMG_EXT | OFFICE_EXT


def _read_pdf(path: str) -> str:
    import fitz
    with fitz.open(path) as doc:
        text = "".join(page.get_text() for page in doc)
        if len(text.strip()) < 20:                       # likely a scanned / image-only PDF
            text += _ocr_pdf_pages(doc)
        return text


def _ocr_pdf_pages(doc) -> str:
    try:
        import pytesseract, io
        from PIL import Image
    except ImportError:
        log.warning("pytesseract/Pillow not installed; skipping OCR for scanned PDF")
        return ""
    out = []
    for page in list(doc)[: config.MAX_OCR_PAGES]:
        try:
            img = Image.open(io.BytesIO(page.get_pixmap(dpi=200).tobytes("png")))
            out.append(pytesseract.image_to_string(img))
        except Exception as e:
            log.warning("OCR failed on a page: %s", e)
    return "\n".join(out)


def _read_docx(path: str) -> str:
    import docx
    d = docx.Document(path)
    parts = [p.text for p in d.paragraphs]
    parts += [cell.text for table in d.tables for row in table.rows for cell in row.cells]
    return "\n".join(parts)


def _read_xlsx(path: str) -> str:
    try:
        import openpyxl
    except ImportError:
        log.warning("openpyxl not installed; skipping %s", path)
        return ""
    wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
    out = []
    for ws in wb.worksheets:
        for row in ws.iter_rows(values_only=True):
            out.append(" ".join(str(c) for c in row if c is not None))
    return "\n".join(out)


def _read_pptx(path: str) -> str:
    try:
        from pptx import Presentation
    except ImportError:
        log.warning("python-pptx not installed; skipping %s", path)
        return ""
    prs = Presentation(path)
    out = []
    for slide in prs.slides:
        for shape in slide.shapes:
            if shape.has_text_frame:
                out.append(shape.text_frame.text)
    return "\n".join(out)


def _read_rtf(path: str) -> str:
    try:
        from striprtf.striprtf import rtf_to_text
    except ImportError:
        log.warning("striprtf not installed; skipping %s", path)
        return ""
    with open(path, errors="ignore") as f:
        return rtf_to_text(f.read())


def _read_image(path: str) -> str:
    try:
        import pytesseract
        from PIL import Image
    except ImportError:
        log.warning("pytesseract/Pillow not installed; skipping image %s", path)
        return ""
    return pytesseract.image_to_string(Image.open(path))


def _read_plain(path: str) -> str:
    with open(path, errors="ignore") as f:
        return f.read()


_READERS = {
    ".pdf": _read_pdf, ".docx": _read_docx, ".xlsx": _read_xlsx,
    ".pptx": _read_pptx, ".rtf": _read_rtf,
}


def extract_text(path: str) -> str:
    """Return the readable text of one file, or '' if the type/content can't be read."""
    ext = os.path.splitext(path)[1].lower()
    try:
        if ext in _READERS:
            return _READERS[ext](path)
        if ext in IMG_EXT:
            return _read_image(path)
        if ext in TEXT_EXT:
            return _read_plain(path)
    except Exception as e:
        log.warning("Could not read %s: %s", path, e)
    return ""
